"""Legacy desktop edition of Bunda Math."""

from __future__ import annotations

import csv
import hashlib
import os
import random
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    from gtts import gTTS
except ImportError:
    gTTS = None

try:
    import pygame
except ImportError:
    pygame = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


APP_NAME = "Bunda Math"
SAMPLE_QUESTIONS = [
    [3, -7, 8, 10, 21], [5, 2, 9, -1, 4], [7, -3, 6, 2, -9],
    [9, 4, -8, 5, 3], [4, -5, 7, -3, 2], [6, 3, -4, 8, -6],
    [8, -2, 5, -7, 1], [2, 7, -3, 4, -5], [1, -6, 9, -1, 8],
    [10, 1, -2, 6, -4],
]
COLORS = {
    "ink": "#17211B", "muted": "#68736C", "paper": "#F5F2E9",
    "card": "#FFFDF7", "line": "#DED9CA", "forest": "#164B35",
    "forest_dark": "#0F3827", "lime": "#C9F269", "lime_dark": "#A9D849",
    "yellow": "#FFD766", "coral": "#F26B52", "blue": "#5C7CFA",
    "white": "#FFFFFF",
}


def parse_number(value):
    """Return an integer from a spreadsheet cell, or None for non-numbers."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace(" ", "")
        if not text:
            return None
        if text.count(",") == 1 and "." not in text:
            text = text.replace(",", ".")
        try:
            number = float(text)
        except ValueError:
            return None
    return int(number) if number.is_integer() else None


def columns_to_questions(rows):
    """Convert rows of cells to question columns while ignoring labels/blanks."""
    width = max((len(row) for row in rows), default=0)
    questions = []
    for column in range(width):
        values = []
        for row in rows:
            if column < len(row):
                number = parse_number(row[column])
                if number is not None:
                    values.append(number)
        if values:
            questions.append(values)
    if not questions:
        raise ValueError("Tidak ada angka yang dapat dibaca di dalam file.")
    return questions


def load_questions(path):
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        with open(path, "r", encoding="utf-8-sig", newline="") as handle:
            sample = handle.read(4096)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.reader(handle, dialect))
    elif suffix in (".xlsx", ".xlsm"):
        if load_workbook is None:
            raise RuntimeError("Dukungan Excel belum terpasang. Jalankan: pip install openpyxl")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = [list(row) for row in workbook.active.iter_rows(values_only=True)]
        finally:
            workbook.close()
    else:
        raise ValueError("Gunakan file CSV atau Excel (.xlsx).")
    return columns_to_questions(rows)


def spoken_number(value, first=False):
    if first:
        return ("negatif " + str(abs(value))) if value < 0 else str(value)
    if value < 0:
        return "dikurangi " + str(abs(value))
    if value > 0:
        return "ditambah " + str(value)
    return "ditambah nol"


def generated_questions(level, count=10):
    settings = {
        "Santai": (4, 1, 9, 0.25),
        "Fokus": (5, 2, 15, 0.40),
        "Juara": (7, 5, 30, 0.50),
    }
    length, low, high, negative_chance = settings[level]
    questions = []
    for _ in range(count):
        values = [random.randint(low, high)]
        for _ in range(length - 1):
            value = random.randint(low, high)
            values.append(-value if random.random() < negative_chance else value)
        questions.append(values)
    return questions


class SpeechService:
    """Generate Indonesian Google speech once and retain it between sessions."""

    def __init__(self, root):
        self.root = root
        local_data = os.environ.get("LOCALAPPDATA", str(Path.home()))
        self.cache_dir = Path(local_data) / "BundaMath" / "speech"
        self.cache_dir.mkdir(parents=True, exist_ok=True)
        self.available = gTTS is not None and pygame is not None
        self.audio_ready = False
        self.cancel_id = 0
        if self.available:
            try:
                pygame.mixer.init(frequency=44100)
                self.audio_ready = True
            except Exception:
                self.available = False

    def path_for(self, text, slow=False):
        key = ("slow:" if slow else "normal:") + text
        digest = hashlib.sha256(key.encode("utf-8")).hexdigest()
        return self.cache_dir / (digest + ".mp3")

    def ensure(self, text, slow=False):
        path = self.path_for(text, slow)
        if not path.exists() or path.stat().st_size == 0:
            temporary = path.with_suffix(".tmp.mp3")
            try:
                gTTS(text=text, lang="id", slow=slow).save(str(temporary))
                temporary.replace(path)
            finally:
                if temporary.exists():
                    temporary.unlink()
        return path

    def prepare(self, phrases, slow, progress, complete):
        unique = sorted(set(phrases))
        if not self.available:
            self.root.after(0, lambda: complete(False))
            return

        def work():
            failures = 0
            done = 0
            with ThreadPoolExecutor(max_workers=4) as executor:
                futures = [executor.submit(self.ensure, phrase, slow) for phrase in unique]
                for future in as_completed(futures):
                    try:
                        future.result()
                    except Exception:
                        failures += 1
                    done += 1
                    self.root.after(0, lambda d=done, t=len(unique): progress(d, t))
            self.root.after(0, lambda: complete(failures == 0))

        threading.Thread(target=work, daemon=True).start()

    def speak(self, text, slow, finished):
        self.stop()
        request_id = self.cancel_id
        if not self.available:
            self.root.after(0, finished)
            return
        try:
            path = self.path_for(text, slow)
            if not path.exists():
                self.root.after(0, finished)
                return
            pygame.mixer.music.load(str(path))
            pygame.mixer.music.play()
        except Exception:
            self.root.after(0, finished)
            return

        def poll():
            if request_id != self.cancel_id:
                return
            if pygame.mixer.music.get_busy():
                self.root.after(40, poll)
            else:
                finished()

        self.root.after(60, poll)

    def stop(self):
        self.cancel_id += 1
        if self.audio_ready:
            try:
                pygame.mixer.music.stop()
            except Exception:
                pass


class BundaMathApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_NAME)
        self.root.configure(bg=COLORS["paper"])
        self.root.geometry("1180x760")
        self.root.minsize(860, 620)
        self.root.protocol("WM_DELETE_WINDOW", self.close)
        self.root.bind("<Return>", self.handle_enter)
        self.root.bind("<space>", self.handle_space)
        self.root.bind("<Escape>", self.handle_escape)

        self.speech = SpeechService(root)
        self.questions = [list(question) for question in SAMPLE_QUESTIONS]
        self.source_name = "Paket contoh"
        self.source_path = None
        self.results = []
        self.round_index = 0
        self.number_index = 0
        self.state = "menu"
        self.timer_id = None
        self.round_started = 0.0
        self.voice_ready = False

        self.level_var = tk.StringVar(value="Fokus")
        self.gap_var = tk.DoubleVar(value=0.45)
        self.voice_var = tk.BooleanVar(value=True)
        self.slow_voice_var = tk.BooleanVar(value=False)
        self.feedback_var = tk.StringVar(value="Langsung")
        self.answer_var = tk.StringVar()
        self.file_var = tk.StringVar()
        self.status_var = tk.StringVar()
        self.question_var = tk.StringVar()
        self.progress_var = tk.StringVar()
        self.hint_var = tk.StringVar()

        self.configure_styles()
        self.show_menu()

    def configure_styles(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("Game.Horizontal.TProgressbar", troughcolor="#2B604A",
                        background=COLORS["lime"], bordercolor=COLORS["forest"],
                        lightcolor=COLORS["lime"], darkcolor=COLORS["lime"])
        style.configure("Menu.Horizontal.TProgressbar", troughcolor=COLORS["line"],
                        background=COLORS["forest"], bordercolor=COLORS["card"],
                        lightcolor=COLORS["forest"], darkcolor=COLORS["forest"])

    def clear(self):
        self.cancel_timer()
        self.speech.stop()
        for widget in self.root.winfo_children():
            widget.destroy()

    def label(self, parent, text="", size=12, weight="normal", color=None, bg=None, **kwargs):
        return tk.Label(parent, text=text, font=("Segoe UI", size, weight),
                        fg=color or COLORS["ink"], bg=bg or parent.cget("bg"), **kwargs)

    def button(self, parent, text, command, bg, fg=None, **kwargs):
        padx = kwargs.pop("padx", 18)
        pady = kwargs.pop("pady", 11)
        button = tk.Button(parent, text=text, command=command, bg=bg,
                           fg=fg or COLORS["white"], activebackground=bg,
                           activeforeground=fg or COLORS["white"], relief="flat",
                           bd=0, cursor="hand2", font=("Segoe UI", 11, "bold"),
                           padx=padx, pady=pady, **kwargs)
        return button

    def show_menu(self):
        self.clear()
        self.state = "menu"
        self.root.attributes("-fullscreen", False)
        self.root.configure(bg=COLORS["paper"])

        shell = tk.Frame(self.root, bg=COLORS["paper"])
        shell.pack(expand=True, fill="both", padx=44, pady=30)
        shell.columnconfigure(0, weight=6)
        shell.columnconfigure(1, weight=5)
        shell.rowconfigure(1, weight=1)

        brand = tk.Frame(shell, bg=COLORS["paper"])
        brand.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 22))
        badge = self.label(brand, "BM", 11, "bold", COLORS["forest"], COLORS["lime"], padx=10, pady=7)
        badge.pack(side="left")
        self.label(brand, "BUNDA MATH", 12, "bold", COLORS["forest"]).pack(side="left", padx=10)
        self.label(brand, "Dengar. Hitung. Hebat.", 10, color=COLORS["muted"]).pack(side="right")

        hero = tk.Frame(shell, bg=COLORS["forest"], padx=42, pady=38)
        hero.grid(row=1, column=0, sticky="nsew", padx=(0, 12))
        hero.rowconfigure(3, weight=1)
        self.label(hero, "LATIHAN ARITMATIKA", 10, "bold", COLORS["lime"], COLORS["forest"]).grid(sticky="w")
        self.label(hero, "Hitung lebih\ncepat, satu suara\npada satu waktu.", 32, "bold",
                   COLORS["white"], COLORS["forest"], justify="left").grid(sticky="w", pady=(16, 12))
        self.label(hero, "Dengarkan rangkaian angka, simpan di ingatan,\nlalu masukkan hasil akhirnya.",
                   12, color="#C7D7CE", bg=COLORS["forest"], justify="left").grid(sticky="w")
        steps = tk.Frame(hero, bg=COLORS["forest"])
        steps.grid(row=3, sticky="sw", pady=(28, 0))
        for number, text in (("01", "DENGARKAN"), ("02", "HITUNG"), ("03", "JAWAB")):
            box = tk.Frame(steps, bg=COLORS["forest"])
            box.pack(side="left", padx=(0, 26))
            self.label(box, number, 18, "bold", COLORS["yellow"], COLORS["forest"]).pack(anchor="w")
            self.label(box, text, 9, "bold", "#C7D7CE", COLORS["forest"]).pack(anchor="w")

        panel = tk.Frame(shell, bg=COLORS["card"], padx=34, pady=30,
                         highlightbackground=COLORS["line"], highlightthickness=1)
        panel.grid(row=1, column=1, sticky="nsew", padx=(12, 0))
        panel.columnconfigure(0, weight=1)
        self.label(panel, "Siapkan permainan", 22, "bold").grid(sticky="w")
        self.label(panel, "Pilih latihan, atur suara, lalu mulai.", 10, color=COLORS["muted"]).grid(sticky="w", pady=(4, 20))

        self.section_title(panel, "SOAL LATIHAN", 2)
        file_row = tk.Frame(panel, bg=COLORS["card"])
        file_row.grid(sticky="ew")
        file_row.columnconfigure(0, weight=1)
        self.file_var.set(f"{self.source_name}  |  {len(self.questions)} soal")
        file_card = tk.Frame(file_row, bg="#F0EDE4", padx=13, pady=11)
        file_card.grid(row=0, column=0, sticky="ew")
        self.label(file_card, textvariable=self.file_var, size=10, weight="bold", bg="#F0EDE4",
                   anchor="w").pack(fill="x")
        self.button(file_row, "Pilih file", self.choose_file, COLORS["forest"], padx=13, pady=9).grid(row=0, column=1, padx=(8, 0))

        random_row = tk.Frame(panel, bg=COLORS["card"])
        random_row.grid(sticky="ew", pady=(10, 0))
        random_row.columnconfigure(0, weight=1)
        self.label(random_row, "Kesulitan acak", 10, "bold").grid(row=0, column=0, sticky="w")
        levels = ttk.Combobox(random_row, textvariable=self.level_var, values=("Santai", "Fokus", "Juara"),
                              state="readonly", width=10, font=("Segoe UI", 10))
        levels.grid(row=0, column=1, padx=(8, 6), ipady=3)
        self.button(random_row, "Buat", self.use_generated, COLORS["yellow"], COLORS["ink"],
                    padx=12, pady=8).grid(row=0, column=2)

        self.section_title(panel, "SUARA & TEMPO", 16)
        voice_row = tk.Frame(panel, bg=COLORS["card"])
        voice_row.grid(sticky="ew")
        tk.Checkbutton(voice_row, text="Suara Indonesia", variable=self.voice_var,
                       bg=COLORS["card"], activebackground=COLORS["card"], selectcolor=COLORS["card"],
                       font=("Segoe UI", 10, "bold"), fg=COLORS["ink"]).pack(side="left")
        tk.Checkbutton(voice_row, text="Tempo bicara pelan", variable=self.slow_voice_var,
                       bg=COLORS["card"], activebackground=COLORS["card"], selectcolor=COLORS["card"],
                       font=("Segoe UI", 10), fg=COLORS["ink"]).pack(side="right")
        gap_row = tk.Frame(panel, bg=COLORS["card"])
        gap_row.grid(sticky="ew", pady=(8, 0))
        self.label(gap_row, "Jeda setelah suara", 10).pack(side="left")
        gap_value = self.label(gap_row, "", 10, "bold", COLORS["forest"])
        gap_value.pack(side="right")
        slider = tk.Scale(panel, from_=0.15, to=1.50, resolution=0.05, orient="horizontal",
                          variable=self.gap_var, bg=COLORS["card"], fg=COLORS["muted"],
                          troughcolor=COLORS["line"], activebackground=COLORS["forest"],
                          highlightthickness=0, showvalue=False, command=lambda value: gap_value.configure(text=f"{float(value):.2f} detik"))
        slider.grid(sticky="ew")
        gap_value.configure(text=f"{self.gap_var.get():.2f} detik")

        self.section_title(panel, "HASIL", 12)
        feedback = tk.Frame(panel, bg=COLORS["card"])
        feedback.grid(sticky="ew")
        for text in ("Langsung", "Di akhir"):
            tk.Radiobutton(feedback, text=text, value=text, variable=self.feedback_var,
                           bg=COLORS["card"], activebackground=COLORS["card"], selectcolor=COLORS["card"],
                           font=("Segoe UI", 10), fg=COLORS["ink"]).pack(side="left", padx=(0, 16))

        self.status_var.set("Siap dimainkan")
        self.label(panel, textvariable=self.status_var, size=9, color=COLORS["muted"], anchor="w").grid(sticky="ew", pady=(14, 7))
        self.start_button = self.button(panel, "MULAI LATIHAN  ->", self.start_game, COLORS["lime"], COLORS["forest_dark"], pady=14)
        self.start_button.grid(sticky="ew")

    def section_title(self, parent, text, top):
        self.label(parent, text, 9, "bold", COLORS["muted"], COLORS["card"]).grid(sticky="w", pady=(top, 7))

    def choose_file(self):
        path = filedialog.askopenfilename(title="Pilih soal", filetypes=[
            ("Soal CSV atau Excel", "*.csv *.xlsx *.xlsm"), ("CSV", "*.csv"), ("Excel", "*.xlsx *.xlsm")])
        if not path:
            return
        try:
            questions = load_questions(path)
        except Exception as error:
            messagebox.showerror("File tidak dapat dibaca", str(error), parent=self.root)
            return
        self.questions = questions
        self.source_path = path
        self.source_name = Path(path).name
        self.file_var.set(f"{self.source_name}  |  {len(self.questions)} soal")
        self.status_var.set(f"Berhasil memuat {len(self.questions)} soal")

    def use_generated(self):
        self.questions = generated_questions(self.level_var.get())
        self.source_path = None
        self.source_name = f"Latihan {self.level_var.get()}"
        self.file_var.set(f"{self.source_name}  |  {len(self.questions)} soal")
        self.status_var.set("Paket soal baru sudah dibuat")

    def phrases_for_game(self):
        phrases = []
        for index, question in enumerate(self.questions, 1):
            phrases.append(f"Soal nomor {index}. Bersiap.")
            phrases.extend(spoken_number(value, position == 0) for position, value in enumerate(question))
            phrases.append(f"Jawabannya {sum(question)}")
        return phrases

    def start_game(self):
        if not self.questions:
            self.status_var.set("Pilih atau buat soal terlebih dahulu")
            return
        self.results = []
        self.round_index = 0
        self.voice_ready = not self.voice_var.get()
        if self.voice_var.get():
            self.show_preparing()
            self.speech.prepare(self.phrases_for_game(), self.slow_voice_var.get(),
                                self.update_preparing, self.preparing_complete)
        else:
            self.show_game()
            self.begin_round()

    def show_preparing(self):
        self.clear()
        self.state = "preparing"
        self.root.configure(bg=COLORS["forest"])
        frame = tk.Frame(self.root, bg=COLORS["forest"])
        frame.pack(expand=True)
        self.label(frame, "MENYIAPKAN SUARA", 10, "bold", COLORS["lime"], COLORS["forest"]).pack()
        self.label(frame, "Sebentar, ya.", 34, "bold", COLORS["white"], COLORS["forest"]).pack(pady=(12, 8))
        self.prepare_status = self.label(frame, "Memeriksa cache suara...", 11, color="#C7D7CE", bg=COLORS["forest"])
        self.prepare_status.pack(pady=(0, 18))
        self.prepare_progress = ttk.Progressbar(frame, style="Game.Horizontal.TProgressbar", length=380, maximum=100)
        self.prepare_progress.pack()

    def update_preparing(self, done, total):
        if self.state != "preparing":
            return
        self.prepare_progress["value"] = done / max(total, 1) * 100
        self.prepare_status.configure(text=f"Menyiapkan suara {done} dari {total}")

    def preparing_complete(self, successful):
        if self.state != "preparing":
            return
        self.voice_ready = successful
        if not successful:
            messagebox.showwarning("Suara tidak tersedia",
                                   "Sebagian suara gagal dibuat. Permainan dilanjutkan tanpa suara.\n\nPeriksa koneksi internet untuk penggunaan pertama.",
                                   parent=self.root)
            self.voice_var.set(False)
        self.show_game()
        self.begin_round()

    def show_game(self):
        self.clear()
        self.state = "ready"
        self.root.configure(bg=COLORS["forest"])
        try:
            self.root.attributes("-fullscreen", True)
        except tk.TclError:
            self.root.state("zoomed")

        self.game = tk.Frame(self.root, bg=COLORS["forest"])
        self.game.pack(expand=True, fill="both", padx=34, pady=24)
        self.game.columnconfigure(0, weight=1)
        self.game.rowconfigure(1, weight=1)

        top = tk.Frame(self.game, bg=COLORS["forest"])
        top.grid(row=0, sticky="ew")
        top.columnconfigure(1, weight=1)
        self.label(top, "BUNDA MATH", 11, "bold", COLORS["lime"], COLORS["forest"]).grid(row=0, column=0, sticky="w")
        self.progress_label = self.label(top, "", 10, "bold", COLORS["white"], COLORS["forest"])
        self.progress_label.grid(row=0, column=2, sticky="e")
        self.progress_bar = ttk.Progressbar(top, style="Game.Horizontal.TProgressbar", maximum=len(self.questions))
        self.progress_bar.grid(row=1, column=0, columnspan=3, sticky="ew", pady=(12, 0))

        self.stage = tk.Frame(self.game, bg=COLORS["forest"])
        self.stage.grid(row=1, sticky="nsew")
        self.stage.columnconfigure(0, weight=1)
        self.stage.rowconfigure(0, weight=1)
        self.display = tk.Frame(self.stage, bg=COLORS["forest"])
        self.display.grid(row=0, sticky="nsew")
        self.display.columnconfigure(0, weight=1)
        self.display.rowconfigure(0, weight=1)

        bottom = tk.Frame(self.game, bg=COLORS["forest"])
        bottom.grid(row=2, sticky="ew")
        self.voice_status = self.label(bottom, "SUARA AKTIF" if self.voice_var.get() else "TANPA SUARA",
                                       9, "bold", COLORS["lime"], COLORS["forest"])
        self.voice_status.pack(side="left")
        self.label(bottom, "Spasi: ulangi  |  Esc: keluar", 9, color="#AFC4B8", bg=COLORS["forest"]).pack(side="right")

    def clear_display(self):
        for widget in self.display.winfo_children():
            widget.destroy()

    def begin_round(self):
        if self.round_index >= len(self.questions):
            self.finish_game()
            return
        self.cancel_timer()
        self.number_index = 0
        self.progress_label.configure(text=f"SOAL {self.round_index + 1} / {len(self.questions)}")
        self.progress_bar["value"] = self.round_index
        self.clear_display()
        self.state = "ready"
        inner = tk.Frame(self.display, bg=COLORS["forest"])
        inner.grid(row=0, column=0)
        self.label(inner, f"SOAL {self.round_index + 1}", 11, "bold", COLORS["yellow"], COLORS["forest"]).pack()
        self.label(inner, "Siap mendengarkan?", 34, "bold", COLORS["white"], COLORS["forest"]).pack(pady=(12, 8))
        self.label(inner, "Angka akan muncul dan dibacakan satu per satu.", 11,
                   color="#C7D7CE", bg=COLORS["forest"]).pack()
        self.round_started = time.monotonic()
        self.play_phrase(f"Soal nomor {self.round_index + 1}. Bersiap.", lambda: self.schedule(350, self.show_next_number))

    def play_phrase(self, phrase, finished):
        if self.voice_var.get() and self.voice_ready:
            self.speech.speak(phrase, self.slow_voice_var.get(), finished)
        else:
            finished()

    def schedule(self, delay, callback):
        self.cancel_timer()
        self.timer_id = self.root.after(delay, callback)

    def cancel_timer(self):
        if self.timer_id is not None:
            try:
                self.root.after_cancel(self.timer_id)
            except tk.TclError:
                pass
            self.timer_id = None

    def show_next_number(self):
        question = self.questions[self.round_index]
        if self.number_index >= len(question):
            self.show_answer_input()
            return
        self.state = "numbers"
        value = question[self.number_index]
        phrase = spoken_number(value, self.number_index == 0)
        self.number_index += 1
        self.clear_display()
        inner = tk.Frame(self.display, bg=COLORS["forest"])
        inner.grid(row=0, column=0)
        self.label(inner, f"ANGKA {self.number_index} DARI {len(question)}", 10, "bold",
                   COLORS["lime"], COLORS["forest"]).pack()
        number_color = COLORS["yellow"] if value < 0 else COLORS["white"]
        text = f"-{abs(value)}" if value < 0 else f"+{value}" if self.number_index > 1 else str(value)
        self.label(inner, text, 82, "bold", number_color, COLORS["forest"]).pack(pady=(4, 0))

        def spoken():
            visual_time = 0 if self.voice_var.get() and self.voice_ready else 800
            self.schedule(int(self.gap_var.get() * 1000) + visual_time, self.show_next_number)

        self.play_phrase(phrase, spoken)

    def show_answer_input(self):
        self.cancel_timer()
        self.state = "answering"
        self.clear_display()
        inner = tk.Frame(self.display, bg=COLORS["forest"])
        inner.grid(row=0, column=0)
        self.label(inner, "SELESAI MENGHITUNG", 10, "bold", COLORS["lime"], COLORS["forest"]).pack()
        self.label(inner, "Berapa hasilnya?", 34, "bold", COLORS["white"], COLORS["forest"]).pack(pady=(10, 18))
        entry_frame = tk.Frame(inner, bg=COLORS["white"], padx=4, pady=4)
        entry_frame.pack()
        self.answer_var.set("")
        self.answer_entry = tk.Entry(entry_frame, textvariable=self.answer_var, width=9, justify="center",
                                     font=("Segoe UI", 30, "bold"), relief="flat", bd=0,
                                     bg=COLORS["white"], fg=COLORS["ink"], insertbackground=COLORS["forest"])
        self.answer_entry.pack(ipady=8)
        self.answer_entry.bind("<space>", self.handle_space)
        self.answer_entry.focus_set()
        self.label(inner, "Ketik jawaban lalu tekan Enter", 10, color="#C7D7CE", bg=COLORS["forest"]).pack(pady=(12, 8))
        self.button(inner, "DENGAR SEKALI LAGI", self.replay_round, COLORS["yellow"], COLORS["forest_dark"],
                    padx=16, pady=9).pack(pady=(8, 0))

    def replay_round(self):
        if self.state not in ("answering", "numbers"):
            return
        self.number_index = 0
        self.show_next_number()

    def submit_answer(self):
        text = self.answer_var.get().strip().replace(" ", "")
        try:
            answer = int(text)
        except ValueError:
            self.answer_entry.configure(bg="#FFE8E3")
            self.root.bell()
            return
        expected = sum(self.questions[self.round_index])
        correct = answer == expected
        self.results.append({
            "question": self.round_index + 1,
            "numbers": list(self.questions[self.round_index]),
            "answer": answer,
            "expected": expected,
            "correct": correct,
            "seconds": round(time.monotonic() - self.round_started, 1),
        })
        if self.feedback_var.get() == "Langsung":
            self.show_round_feedback(correct, expected)
        else:
            self.advance_round()

    def show_round_feedback(self, correct, expected):
        self.state = "feedback"
        self.clear_display()
        inner = tk.Frame(self.display, bg=COLORS["forest"])
        inner.grid(row=0, column=0)
        if correct:
            self.label(inner, "TEPAT!", 11, "bold", COLORS["lime"], COLORS["forest"]).pack()
            self.label(inner, "Kerja bagus.", 38, "bold", COLORS["white"], COLORS["forest"]).pack(pady=(10, 8))
        else:
            self.label(inner, "BELUM TEPAT", 11, "bold", COLORS["coral"], COLORS["forest"]).pack()
            self.label(inner, str(expected), 62, "bold", COLORS["yellow"], COLORS["forest"]).pack(pady=(4, 0))
            self.label(inner, "adalah jawaban yang benar", 11, color="#C7D7CE", bg=COLORS["forest"]).pack()
        self.label(inner, "Tekan Enter untuk soal berikutnya", 10, color="#C7D7CE", bg=COLORS["forest"]).pack(pady=(24, 0))
        self.play_phrase(f"Jawabannya {expected}", lambda: None)

    def advance_round(self):
        self.round_index += 1
        self.begin_round()

    def finish_game(self):
        self.clear()
        self.state = "finished"
        self.root.attributes("-fullscreen", False)
        self.root.configure(bg=COLORS["paper"])
        correct = sum(result["correct"] for result in self.results)
        total = len(self.results)
        percent = round(correct / total * 100) if total else 0

        shell = tk.Frame(self.root, bg=COLORS["paper"])
        shell.pack(expand=True, fill="both", padx=48, pady=32)
        shell.columnconfigure(0, weight=1)
        shell.rowconfigure(1, weight=1)
        header = tk.Frame(shell, bg=COLORS["paper"])
        header.grid(row=0, sticky="ew", pady=(0, 20))
        self.label(header, "LATIHAN SELESAI", 10, "bold", COLORS["forest"]).pack(anchor="w")
        self.label(header, "Lihat kemajuanmu.", 28, "bold").pack(anchor="w", pady=(5, 0))

        content = tk.Frame(shell, bg=COLORS["paper"])
        content.grid(row=1, sticky="nsew")
        content.columnconfigure(1, weight=1)
        content.rowconfigure(0, weight=1)

        score = tk.Frame(content, bg=COLORS["forest"], padx=34, pady=30, width=290)
        score.grid(row=0, column=0, sticky="nsew", padx=(0, 14))
        score.grid_propagate(False)
        self.label(score, "SKOR HARI INI", 10, "bold", COLORS["lime"], COLORS["forest"]).pack(anchor="w")
        self.label(score, f"{percent}%", 48, "bold", COLORS["white"], COLORS["forest"]).pack(anchor="w", pady=(16, 0))
        self.label(score, f"{correct} dari {total} jawaban benar", 11, color="#C7D7CE", bg=COLORS["forest"]).pack(anchor="w")
        ttk.Progressbar(score, style="Game.Horizontal.TProgressbar", maximum=max(total, 1),
                        value=correct, length=220).pack(anchor="w", pady=(22, 26), fill="x")
        note = "Luar biasa! Pertahankan." if percent >= 80 else "Sedikit lagi. Coba ulangi latihan ini."
        self.label(score, note, 12, "bold", COLORS["yellow"], COLORS["forest"],
                   wraplength=210, justify="left").pack(anchor="w")

        review = tk.Frame(content, bg=COLORS["card"], padx=24, pady=22,
                          highlightbackground=COLORS["line"], highlightthickness=1)
        review.grid(row=0, column=1, sticky="nsew", padx=(14, 0))
        review.columnconfigure(0, weight=1)
        review.rowconfigure(1, weight=1)
        self.label(review, "Rincian jawaban", 17, "bold").grid(row=0, sticky="w", pady=(0, 12))
        table_frame = tk.Frame(review, bg=COLORS["card"])
        table_frame.grid(row=1, sticky="nsew")
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)
        columns = ("question", "sequence", "answer", "correct")
        table = ttk.Treeview(table_frame, columns=columns, show="headings", height=10)
        table.heading("question", text="Soal")
        table.heading("sequence", text="Rangkaian")
        table.heading("answer", text="Jawabanmu")
        table.heading("correct", text="Hasil")
        table.column("question", width=55, anchor="center", stretch=False)
        table.column("sequence", width=280, anchor="w")
        table.column("answer", width=100, anchor="center", stretch=False)
        table.column("correct", width=110, anchor="center", stretch=False)
        scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=table.yview)
        table.configure(yscrollcommand=scrollbar.set)
        table.grid(row=0, column=0, sticky="nsew")
        scrollbar.grid(row=0, column=1, sticky="ns")
        for result in self.results:
            sequence = "  ".join((f"{value:+d}" if index else str(value))
                                for index, value in enumerate(result["numbers"]))
            status = "Benar" if result["correct"] else f"Benar: {result['expected']}"
            table.insert("", "end", values=(result["question"], sequence, result["answer"], status))

        actions = tk.Frame(shell, bg=COLORS["paper"])
        actions.grid(row=2, sticky="ew", pady=(20, 0))
        self.button(actions, "SIMPAN HASIL", self.save_results, COLORS["yellow"], COLORS["forest_dark"]).pack(side="left")
        self.button(actions, "KEMBALI KE MENU", self.show_menu, COLORS["forest"]).pack(side="right", padx=(10, 0))
        self.button(actions, "MAIN LAGI", self.play_again, COLORS["lime"], COLORS["forest_dark"]).pack(side="right")

    def play_again(self):
        self.results = []
        self.round_index = 0
        self.show_game()
        self.begin_round()

    def save_results(self):
        default = (Path(self.source_path).stem if self.source_path else "latihan") + "_hasil.csv"
        path = filedialog.asksaveasfilename(title="Simpan hasil", initialfile=default,
                                            defaultextension=".csv", filetypes=[("CSV", "*.csv")])
        if not path:
            return
        try:
            with open(path, "w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(["Soal", "Rangkaian", "Jawaban", "Kunci", "Benar", "Detik"])
                for result in self.results:
                    writer.writerow([result["question"], " ".join(map(str, result["numbers"])),
                                     result["answer"], result["expected"],
                                     "Ya" if result["correct"] else "Tidak", result["seconds"]])
            messagebox.showinfo("Hasil tersimpan", f"Hasil disimpan ke:\n{path}", parent=self.root)
        except OSError as error:
            messagebox.showerror("Gagal menyimpan", str(error), parent=self.root)

    def handle_enter(self, event=None):
        if self.state == "answering":
            self.submit_answer()
            return "break"
        if self.state == "feedback":
            self.advance_round()
            return "break"
        return None

    def handle_space(self, event=None):
        if self.state == "answering":
            self.replay_round()
            return "break"
        return None

    def handle_escape(self, event=None):
        if self.state in ("ready", "numbers", "answering", "feedback", "preparing"):
            if messagebox.askyesno("Keluar dari latihan", "Kembali ke menu? Progres saat ini akan hilang.", parent=self.root):
                self.show_menu()
        return "break"

    def close(self):
        self.cancel_timer()
        self.speech.stop()
        self.root.destroy()


def main():
    root = tk.Tk()
    BundaMathApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
